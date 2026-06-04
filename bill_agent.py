"""
EnergyBae Solar Load Calculator Agent
Extracts data from MSEDCL electricity bills and fills the Excel template.
"""

import anthropic
import base64
import json
import shutil
import sys
from pathlib import Path
from openpyxl import load_workbook


# ── Anthropic client ──────────────────────────────────────────────────────────
client = anthropic.Anthropic()  # reads ANTHROPIC_API_KEY from env


# ── Step 1: Extract bill data using Claude Vision ─────────────────────────────
def extract_bill_data(bill_path: str) -> dict:
    """Send bill (PDF or image) to Claude and extract structured data."""

    bill_path = Path(bill_path)
    suffix = bill_path.suffix.lower()

    # Determine media type
    media_type_map = {
        ".pdf": "application/pdf",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".webp": "image/webp",
    }
    media_type = media_type_map.get(suffix)
    if not media_type:
        raise ValueError(f"Unsupported file type: {suffix}")

    # Read and encode file
    with open(bill_path, "rb") as f:
        file_data = base64.standard_b64encode(f.read()).decode("utf-8")

    # Build content block
    if media_type == "application/pdf":
        file_block = {
            "type": "document",
            "source": {"type": "base64", "media_type": media_type, "data": file_data},
        }
    else:
        file_block = {
            "type": "image",
            "source": {"type": "base64", "media_type": media_type, "data": file_data},
        }

    prompt = """You are an expert at reading Indian MSEDCL electricity bills.
Extract the following fields from this bill and return ONLY a valid JSON object — no explanation, no markdown, no backticks.

Required fields:
{
  "consumer_name": "full name as printed on bill",
  "consumer_no": "consumer/account number",
  "fixed_charges": <number in rupees, e.g. 130.05>,
  "sanctioned_load": <number in kW, e.g. 7.0>,
  "monthly_data": [
    {"month": "Mon-YYYY", "units": <integer>, "bill_amount": <number or null>},
    ...up to 12 months in reverse chronological order (latest first)
  ]
}

Rules:
- monthly_data must have at least 1 entry, ideally 11-12 months
- month format: "May-2026", "April-2026", etc.
- bill_amount: include only for the most recent month if clearly shown, else null
- If a field is not found, use null
- Return ONLY the JSON, nothing else"""

    response = client.messages.create(
        model="claude-opus-4-5",
        max_tokens=1500,
        messages=[
            {
                "role": "user",
                "content": [file_block, {"type": "text", "text": prompt}],
            }
        ],
    )

    raw = response.content[0].text.strip()
    # Strip any accidental markdown fences
    if raw.startswith("```"):
        raw = raw.split("```")[1]
        if raw.startswith("json"):
            raw = raw[4:]
    return json.loads(raw.strip())


# ── Step 2: Fill Excel template ───────────────────────────────────────────────
def fill_excel_template(data: dict, template_path: str, output_path: str):
    """Fill only the input cells in the Excel template."""

    shutil.copy(template_path, output_path)
    wb = load_workbook(output_path)
    ws = wb.active

    # Header fields
    if data.get("consumer_name"):
        ws["D3"] = data["consumer_name"]
    if data.get("consumer_no"):
        ws["D4"] = data["consumer_no"]
    if data.get("fixed_charges") is not None:
        ws["D5"] = float(data["fixed_charges"])
    if data.get("sanctioned_load") is not None:
        ws["D6"] = float(data["sanctioned_load"])

    # Monthly data — rows 9 to 19 (max 11 months)
    monthly = data.get("monthly_data", [])[:11]
    row_start = 9
    for i, entry in enumerate(monthly):
        row = row_start + i
        if entry.get("month"):
            ws[f"C{row}"] = entry["month"]
        if entry.get("units") is not None:
            ws[f"D{row}"] = int(entry["units"])
        # Bill amount only for first row
        if i == 0 and entry.get("bill_amount") is not None:
            ws[f"E{row}"] = float(entry["bill_amount"])

    # Clear unused rows (in case template had old data)
    for row in range(row_start + len(monthly), 20):
        ws[f"C{row}"] = None
        ws[f"D{row}"] = None
        ws[f"E{row}"] = None

    wb.save(output_path)
    print(f"✅ Excel filled and saved: {output_path}")


# ── Step 3: Print summary ─────────────────────────────────────────────────────
def print_summary(data: dict):
    print("\n" + "="*50)
    print("📋 EXTRACTED BILL DATA")
    print("="*50)
    print(f"  Consumer Name    : {data.get('consumer_name', 'N/A')}")
    print(f"  Consumer No      : {data.get('consumer_no', 'N/A')}")
    print(f"  Fixed Charges    : ₹{data.get('fixed_charges', 'N/A')}")
    print(f"  Sanctioned Load  : {data.get('sanctioned_load', 'N/A')} kW")
    monthly = data.get("monthly_data", [])
    print(f"\n  Monthly Data ({len(monthly)} months):")
    for m in monthly:
        amt = f"  ₹{m['bill_amount']}" if m.get("bill_amount") else ""
        print(f"    {m.get('month','?'):15s} → {m.get('units','?')} units{amt}")
    print("="*50 + "\n")


# ── Main ──────────────────────────────────────────────────────────────────────
def run(bill_path: str, template_path: str = "template.xlsx", output_path: str = None):
    if output_path is None:
        stem = Path(bill_path).stem
        output_path = f"solar_analysis_{stem}.xlsx"

    print(f"🔍 Reading bill: {bill_path}")
    data = extract_bill_data(bill_path)
    print_summary(data)

    print(f"📊 Filling Excel template...")
    fill_excel_template(data, template_path, output_path)

    return output_path, data


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python bill_agent.py <path_to_bill.pdf_or_image>")
        sys.exit(1)

    bill_file = sys.argv[1]
    template_file = sys.argv[2] if len(sys.argv) > 2 else "template.xlsx"
    out_file = sys.argv[3] if len(sys.argv) > 3 else None

    result_path, extracted = run(bill_file, template_file, out_file)
    print(f"🎉 Done! Output saved to: {result_path}")
